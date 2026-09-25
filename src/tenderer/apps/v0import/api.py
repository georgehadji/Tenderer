"""Import one v0 workbook copy (discovery/concierge-phase-b-workbook.xlsx) into v1, with a reconciliation report.

The workbook stays the operator's entry sheet for invitations, routes, client costs and offers (flow F1). This
import records them in v1 and recomputes everything with v1's code: go/no-go per route, the offer check, the
participation guarantee and each document's status. Each result is compared with the value the spreadsheet cached
when it was last saved. Any difference rolls the whole import back unless the operator accepts it (build-plan M8:
"identical numbers"). Metadata only: the workbook holds no files, names or document content (AD5).
"""

import datetime as dt
from dataclasses import dataclass, field, replace
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from django.db import transaction
from django.utils import timezone

from tenderer.apps.alerts import api as alerts
from tenderer.apps.alerts.models import Deadline
from tenderer.apps.audit import api as audit
from tenderer.apps.engagements import api as engagements
from tenderer.apps.engagements.models import Bid, Client, DocumentRecord, Engagement, Resource, packs
from tenderer.core.catalog.tender import Anchor, Tender
from tenderer.core.pricing.money import Money
from tenderer.core.pricing.offer import (
    GuaranteeCheck,
    OfferLine,
    Route,
    check_participation_guarantee,
    guarantee_amount,
    validate_offer,
)
from tenderer.core.rules.checklist import DocumentFact, Item, Status, evaluate
from tenderer.core.rules.deadlines import offer_validity_end

ROWS = range(6, 21)  # routes and offer lines
DOC_ROWS = range(6, 30)
NO_DATA = "λείπουν στοιχεία"
COSTS = {"fuel_l_per_100km": "fuel_cons", "fuel_price_per_l": "fuel_price", "wear_per_km": "wear",
         "opportunity_per_hour": "opp_rate", "extra_insurance_per_year": "ins_year", "bank_rate_per_year": "bank_rate",
         "bank_fee_per_guarantee": "bank_fee", "paid_share": "paid_share", "fuel_increase": "fuel_scn"}
STAGE_DATE = {"Β4": "deadline", "Β5": "deadline", "Β9": "award_docs_date", "Β10": "sign_date"}
CENT = Decimal("0.005")


class WorkbookError(Exception):
    """The file cannot be imported as it is (not a v0 workbook, or never saved by a spreadsheet program)."""


class Differences(Exception):
    def __init__(self, report: "Report") -> None:
        self.report = report
        super().__init__(f"{report.differences} difference(s) between v0 and v1")


@dataclass
class Report:
    lines: list[str] = field(default_factory=list)
    differences: int = 0

    def same(self, what: str) -> None:
        self.lines.append(f"  same  {what}")

    def diff(self, what: str) -> None:
        self.lines.append(f"  DIFF  {what}")
        self.differences += 1

    def note(self, what: str) -> None:
        self.lines.append(f"  note  {what}")

    def compare(self, what: str, v0: object, v1: object) -> None:
        equal = abs(v0 - v1) < CENT if isinstance(v0, Decimal) and isinstance(v1, Decimal) else v0 == v1
        (self.same if equal else self.diff)(f"{what}: v0 {v0}, v1 {v1}")

    def text(self) -> str:
        return "\n".join([*self.lines, f"differences: {self.differences}"])


def import_workbook(path: Path, client: Client, tender: Tender, accept_differences: bool = False,
                    today: dt.date | None = None) -> Report:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)  # cached values; openpyxl parses XML with defusedxml
    except Exception as e:
        raise WorkbookError(f"{path.name}: not a readable .xlsx ({type(e).__name__})") from e
    for sheet in ("Πρόσκληση", "Δρομολόγια", "Προσφορά", "Έγγραφα"):
        if sheet not in wb.sheetnames:
            raise WorkbookError(f"{path.name}: no sheet {sheet!r}; is this the v0 workbook?")
    report = Report()
    with transaction.atomic():
        _import(wb, client, tender, report, today or timezone.localdate())
        if report.differences and not accept_differences:
            raise Differences(report)  # rolls back every record written above
    return report


def _import(wb: Any, client: Client, tender: Tender, report: Report, today: dt.date) -> None:
    value = _named(wb)
    invitation = _text(value("invitation_ref"))
    if not invitation:
        raise WorkbookError("the invitation (sheet Πρόσκληση, B5) is empty")
    dates = {name: _date(value(name)) for name in ("inv_date", "deadline", "sign_date", "contract_end",
                                                    "award_docs_date")}
    report.note(f"client {client.pk}, tender {tender.id} ({tender.version}), invitation {invitation}")
    engagement = (Engagement.objects.filter(client=client, tender_id=tender.id).first()
                  or engagements.open_engagement(client, tender))
    bid = (Bid.objects.filter(engagement=engagement, invitation_ref=invitation).first()
           or engagements.open_bid(engagement, invitation))

    routes = _routes(wb, value, dates, tender, bid, report)
    checks = _offer(wb, routes, bid, report)
    guarantee = _guarantee(wb, routes, checks, dates, tender, report)
    _documents(wb, client, tender, bid, dates, checks, guarantee, report)
    _deadlines(engagement, invitation, dates, today, report)
    audit.record("v0.imported", bid, {"differences": report.differences})


def _routes(wb: Any, value: Any, dates: dict[str, dt.date | None], tender: Tender, bid: Bid,
            report: Report) -> dict[str, Route]:
    ws = wb["Δρομολόγια"]
    sector = packs().sectors[tender.sector]
    deduction = packs().jurisdictions[tender.jurisdiction].payment_deduction_rate
    client_costs = {k: _number(value(name)) for k, name in COSTS.items()}
    contract = {"school_years": _number(value("years")), "signed_on": dates["sign_date"],
                "ends_on": dates["contract_end"]}
    routes: dict[str, Route] = {}
    for r in ROWS:
        code = _text(ws[f"C{r}"].value)
        if not code:
            continue
        escort = _text(ws[f"E{r}"].value) == "ΝΑΙ"
        inputs = {"reference": _number(ws[f"F{r}"].value), "days": _number(ws[f"G{r}"].value),
                  "budget": _number(ws[f"H{r}"].value), "km_per_day": _number(ws[f"J{r}"].value),
                  "hours_per_day": _number(ws[f"K{r}"].value), "escort": escort,
                  "escort_per_day": _number(ws[f"L{r}"].value)}
        if inputs["reference"] and inputs["budget"] and inputs["days"]:
            routes[code] = Route(code, Money(Decimal(inputs["reference"])), Money(Decimal(inputs["budget"])),
                                 int(Decimal(inputs["days"])), group=_text(ws[f"D{r}"].value),
                                 client_said_yes=_text(ws[f"N{r}"].value) == "ΝΑΙ")
            report.compare(f"route {code} participation guarantee (€)", _decimal(ws[f"Z{r}"].value),
                           guarantee_amount(routes[code].budget, tender.offer.participation_guarantee_rate).amount)
        result, warnings = sector.gonogo(inputs, client_costs, contract, tender.offer, deduction)
        v0_net, v0_limit = ws[f"T{r}"].value, ws[f"V{r}"].value
        if not hasattr(result, "break_even"):
            report.compare(f"route {code} inputs", _text(v0_net) or None,
                           NO_DATA if v0_net == NO_DATA else f"missing {', '.join(result.fields)}")
            continue
        if v0_net is None:
            raise WorkbookError("no cached results: open the file in a spreadsheet program, save it, import again")
        engagements.record_gonogo(bid, code, result, warnings)
        report.compare(f"route {code} net per day at 0% (€)", _decimal(v0_net), result.net(0))
        report.compare(f"route {code} largest discount without loss (%)",
                       None if isinstance(v0_limit, str) else int(v0_limit), result.break_even)
    return routes


def _offer(wb: Any, routes: dict[str, Route], bid: Bid, report: Report) -> list[Any]:
    ws = wb["Προσφορά"]
    lines, v0 = [], []
    for r in ROWS:
        code = _text(ws[f"B{r}"].value)
        if not code:
            continue
        lines.append(OfferLine(code, _number(ws[f"F{r}"].value), _money(ws[f"H{r}"].value),
                               _money(ws[f"I{r}"].value), _number(ws[f"J{r}"].value)))
        v0.append((_decimal(ws[f"G{r}"].value), _text(ws[f"M{r}"].value)))
    if not lines:
        report.note("no offer lines yet")
        return []
    checks = list(validate_offer(lines, routes))
    engagements.record_offer_check(bid, checks)
    for check, (v0_price, v0_state) in zip(checks, v0, strict=True):
        code = check.line.route_code
        report.compare(f"offer {code} price (€)", v0_price, check.price.amount if check.price else None)
        report.compare(f"offer {code} has an error", v0_state.startswith("ΛΑΘΟΣ"), check.error is not None)
    return checks


def _guarantee(wb: Any, routes: dict[str, Route], checks: list[Any], dates: dict[str, dt.date | None],
               tender: Tender, report: Report) -> GuaranteeCheck | None:
    if not checks:
        return None
    ws = wb["Προσφορά"]
    offered = [routes[c.line.route_code].budget.amount for c in checks if c.line.route_code in routes]
    minimum = guarantee_amount(Money(sum(offered, Decimal(0))), tender.offer.participation_guarantee_rate)
    kind = _text(ws["E28"].value)
    required = (offer_validity_end(tender.offer, dates["deadline"])
                + dt.timedelta(days=tender.offer.participation_guarantee_extra_days)) if dates["deadline"] else None
    result = check_participation_guarantee(minimum, _money(ws["E26"].value), _date(ws["E27"].value),
                                           {"ΕΝΤΥΠΗ": True, "ΗΛΕΚΤΡΟΝΙΚΗ": False}.get(kind), required)
    v0 = _text(ws["E29"].value)
    v0_ok = v0.startswith("OK") if v0 else None
    report.compare("participation guarantee accepted", v0_ok,
                   result in (GuaranteeCheck.OK, GuaranteeCheck.OK_PAPER) if v0 else None)
    return result


def _documents(wb: Any, client: Client, tender: Tender, bid: Bid, dates: dict[str, dt.date | None],
               checks: list[Any], guarantee: GuaranteeCheck | None, report: Report) -> None:
    ws = wb["Έγγραφα"]
    resources = {r.label: r for r in Resource.objects.filter(client=client)}
    ids = {r.id for r in tender.requirements}
    offer_stage: list[Item] = []
    for r in DOC_ROWS:
        rid, stage, label = _text(ws[f"B{r}"].value), _text(ws[f"A{r}"].value), _text(ws[f"D{r}"].value)
        if rid not in ids:
            report.diff(f"row {r}: {rid!r} is not a requirement of {tender.id}")
            continue
        if label and label not in resources:
            report.diff(f"row {r} {rid}: no resource labelled {label!r}; add it for this client and import again")
            continue
        req = tender.requirement(rid)
        applicable = _text(ws[f"E{r}"].value) != "ΟΧΙ"
        issued, valid = _date(ws[f"G{r}"].value), _date(ws[f"H{r}"].value)
        manual = _text(ws[f"I{r}"].value)
        manual_status = Status(manual) if manual in ("SATISFIED", "NOT_SATISFIED") else None
        if rid == "R11" and checks:  # derived in v1 exactly as in v0: the offer check
            manual_status = Status.NOT_SATISFIED if any(c.error for c in checks) else Status.SATISFIED
        if rid == "R8":  # the guarantee's own expiry, from the Offer sheet
            valid = _date(wb["Προσφορά"]["E27"].value)
        if rid != "R11" and (issued or valid or manual_status or not applicable):  # R11 lives in the bid
            DocumentRecord.objects.update_or_create(
                client=client, resource=resources.get(label), doc_type=req.doc_type,
                defaults={"issued_on": issued, "valid_until": valid, "applicable": applicable,
                          "manual_status": manual_status.value if manual_status else ""})
        submission = dates.get(STAGE_DATE.get(stage, ""))
        key_dates = {k: v for k, v in {
            Anchor.SUBMISSION: submission, Anchor.INVITATION_SENT: dates["inv_date"],
            Anchor.CONTRACT_END: dates["contract_end"],
            Anchor.OFFER_VALIDITY_END: (offer_validity_end(tender.offer, dates["deadline"])
                                        if dates["deadline"] else None),
        }.items() if v is not None}
        [item] = evaluate([req], [DocumentFact(req.doc_type, label, issued, valid, manual_status, applicable)],
                          key_dates)
        if rid == "R8" and guarantee in (GuaranteeCheck.BELOW_MINIMUM, GuaranteeCheck.MISSING_DATA):
            # the date rule alone would pass a guarantee that is too small (§4.3.1.2)
            status = Status.NOT_SATISFIED if guarantee is GuaranteeCheck.BELOW_MINIMUM else Status.UNKNOWN
            item = replace(item, status=status, reason=f"guarantee_{guarantee.value}")
        v0 = _text(ws[f"J{r}"].value)
        report.compare(f"document {rid} {label or 'client'} ({stage})",
                       "NOT_APPLICABLE" if v0 == "ΔΕΝ ΙΣΧΥΕΙ" else v0, item.status.value)
        if stage in ("Β4", "Β5"):
            offer_stage.append(item)
    if offer_stage:
        engagements.record_checklist(bid, offer_stage)


def _deadlines(engagement: Engagement, invitation: str, dates: dict[str, dt.date | None], today: dt.date,
               report: Report) -> None:
    for name, step, section in (("deadline", f"Υποβολή προσφοράς, πρόσκληση {invitation}", ""),
                                ("award_docs_date", "Δικαιολογητικά προσωρινού αναδόχου", "§5.3.1")):
        due = dates[name]
        if due is None or due < today:
            continue
        if Deadline.objects.filter(engagement=engagement, step=step, due_on=due, closed_at__isnull=True).exists():
            continue
        now = timezone.now() if today == timezone.localdate() else dt.datetime.combine(
            today, dt.time(8), tzinfo=timezone.get_current_timezone())
        alerts.add_deadline(Deadline(engagement=engagement, step=step, due_on=due, source_section=section), now)
        report.note(f"deadline {due:%d/%m/%Y}: {step} (calendar e-mail and reminders queued)")


def _named(wb: Any) -> Any:
    def value(name: str) -> object:
        sheet, cell = next(iter(wb.defined_names[name].destinations))
        return wb[sheet][cell.replace("$", "")].value
    return value


def _text(v: object) -> str:
    return "" if v is None else str(v).strip()


def _number(v: object) -> str | None:
    """Spreadsheet numbers arrive as floats; their shortest repr is what the cell shows (never Decimal(float))."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    return str(v) if isinstance(v, int | float) else _text(v)


def _decimal(v: object) -> Decimal | None:
    n = _number(v)
    try:
        return Decimal(n) if n is not None else None
    except ArithmeticError:
        return None


def _money(v: object) -> Money | None:
    d = _decimal(v)
    return Money(d) if d is not None else None


def _date(v: object) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    return v if isinstance(v, dt.date) else None
