"""apps/v0import against PostgreSQL: a filled copy of the v0 workbook imports with zero differences when its cached
results are those of the independent oracle that verified the workbook; one wrong cached value rolls everything
back; a file never saved by a spreadsheet program is refused (build-plan M8). Synthetic data only."""

import datetime as dt
from decimal import Decimal as D
from pathlib import Path

import openpyxl
import pytest
from django.core.management import CommandError, call_command

from tenderer.apps.alerts.models import Deadline, Outbox
from tenderer.apps.engagements.models import Bid, Client, DocumentRecord, Engagement, Resource
from tenderer.apps.v0import import api
from tests.test_taxi_costs import ROUTES, C, o_amt_p, o_breakeven, o_net, o_price

pytestmark = pytest.mark.django_db

TEMPLATE = Path(__file__).resolve().parents[1] / "discovery" / "concierge-phase-b-workbook.xlsx"
INVITATION, DEADLINE, AWARD_DOCS = dt.date(2026, 8, 7), dt.date(2026, 8, 19), dt.date(2026, 9, 10)
TODAY = dt.date(2026, 8, 10)


@pytest.fixture
def client_():
    client = Client.objects.create(name="Πελάτης Δοκιμής", tax_id="123456783", email="client@example.com")
    for kind, label, attrs in (
        ("vehicle", "όχημα 1", {"plate": "ΝΒΚ-1234", "seats": 4, "base_municipality": "Θεσσαλονίκη"}),
        ("driver", "οδηγός 1", {"display_name": "Οδηγός Δοκιμής"}),
        ("escort", "συνοδός 1", {"display_label": "Συνοδός Δοκιμής"}),
    ):
        Resource.objects.create(client=client, sector="taxi_student_transport", kind=kind, label=label,
                                attributes=attrs)
    return client


def filled_workbook(path: Path, cached: bool = True) -> Path:
    """The v0 template as the operator fills it; `cached` puts the oracle's results where the spreadsheet caches
    its own when it saves the file."""
    wb = openpyxl.load_workbook(TEMPLATE)
    inv = wb["Πρόσκληση"]
    for cell, value in {"B4": "Π01", "B5": "ΑΔΑΜ-TEST-1", "B6": INVITATION, "B7": DEADLINE, "B9": C["years"],
                        "B10": C["sign"], "B11": C["end"], "B12": AWARD_DOCS, "B22": float(C["cons"]),
                        "B23": float(C["fuel"]), "B24": float(C["wear"]), "B25": float(C["opp"]),
                        "B26": float(C["ins"]), "B27": float(C["rate"]), "B28": float(C["fee"]),
                        "B29": float(C["share"]), "B30": float(C["scn"])}.items():
        inv[cell] = value
    ws = wb["Δρομολόγια"]
    for row, rt in enumerate(ROUTES, start=6):
        for col, value in {"C": rt["C_"], "E": rt["E"], "F": float(rt["F"]), "G": rt["days"], "H": float(rt["H"]),
                           "J": float(rt["J"]), "K": float(rt["K"]), "L": float(rt["L"]) if rt["L"] else None,
                           "M": rt["M"], "N": "ΝΑΙ" if rt["M"] is not None else "ΟΧΙ"}.items():
            ws[f"{col}{row}"] = value
        if cached:
            be = o_breakeven(rt)
            ws[f"T{row}"] = float(o_net(rt, 0))
            ws[f"V{row}"] = "καμία: ζημιά και με 0%" if be is None else be
            ws[f"Z{row}"] = float(o_amt_p(rt["H"]))
    offer = wb["Προσφορά"]
    route = ROUTES[0]  # G26-0703-Τ2, the client decided 3%
    price = float(o_price(route["F"], 3))
    for cell, value in {"B6": route["C_"], "F6": 3, "H6": price, "I6": price, "J6": 1,
                        "E26": float(o_amt_p(route["H"])), "E27": dt.date(2027, 10, 31), "E28": "ΗΛΕΚΤΡΟΝΙΚΗ"}.items():
        offer[cell] = value
    if cached:
        offer["G6"], offer["M6"], offer["E29"] = price, "OK", "OK"
    docs = wb["Έγγραφα"]
    docs["H6"] = dt.date(2027, 12, 31)  # R1 signature certificate valid past the deadline
    docs["G8"] = dt.date(2026, 8, 10)  # R9 signed after the invitation, before the deadline
    docs["E9"] = docs["E10"] = "ΟΧΙ"  # R10, R24: sole owner of the vehicle
    if cached:
        expected = {6: "SATISFIED", 7: "SATISFIED", 8: "SATISFIED", 9: "ΔΕΝ ΙΣΧΥΕΙ", 10: "ΔΕΝ ΙΣΧΥΕΙ",
                    11: "SATISFIED"}
        for row in range(6, 30):
            docs[f"J{row}"] = expected.get(row, "UNKNOWN")
    out = path / "Π01-ΑΔΑΜ-TEST-1.xlsx"
    wb.save(out)
    return out


def test_import_reproduces_every_v0_number(client_, tender, tmp_path):
    try:
        report = api.import_workbook(filled_workbook(tmp_path), client_, tender, today=TODAY)
    except api.Differences as e:
        pytest.fail(e.report.text())
    assert report.differences == 0
    bid = Bid.objects.get(invitation_ref="ΑΔΑΜ-TEST-1")
    assert set(bid.gonogo_snapshot) == {rt["C_"] for rt in ROUTES}
    for rt in ROUTES:
        assert D(bid.gonogo_snapshot[rt["C_"]]["net_at_0"]).quantize(D("0.0001")) == o_net(rt, 0).quantize(
            D("0.0001"))
        assert bid.gonogo_snapshot[rt["C_"]]["break_even"] == o_breakeven(rt)
    assert [line["error"] for line in bid.offer_check] == [None]
    assert {i["requirement_id"]: i["status"] for i in bid.checklist_snapshot} == {
        "R1": "SATISFIED", "R8": "SATISFIED", "R9": "SATISFIED", "R10": "NOT_APPLICABLE", "R24": "NOT_APPLICABLE",
        "R11": "SATISFIED"}
    assert DocumentRecord.objects.filter(client=client_).count() == 5  # R1, R8, R9, R10, R24: rows with data
    assert sorted(Deadline.objects.values_list("due_on", flat=True)) == [DEADLINE, AWARD_DOCS]
    assert Outbox.objects.count() > 0
    # importing the same file again changes nothing
    again = api.import_workbook(filled_workbook(tmp_path), client_, tender, today=TODAY)
    assert again.differences == 0
    assert (Engagement.objects.count(), Bid.objects.count(), Deadline.objects.count()) == (1, 1, 2)


def test_one_different_number_rolls_everything_back(client_, tender, tmp_path):
    path = filled_workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Δρομολόγια"]["T6"] = wb["Δρομολόγια"]["T6"].value + 1  # v0 says one euro more
    wb.save(path)
    with pytest.raises(api.Differences) as refused:
        api.import_workbook(path, client_, tender, today=TODAY)
    assert refused.value.report.differences == 1
    assert "DIFF  route G26-0703-Τ2 net per day at 0%" in refused.value.report.text()
    assert not Engagement.objects.exists() and not Deadline.objects.exists()


def test_a_file_never_saved_by_a_spreadsheet_is_refused(client_, tender, tmp_path):
    with pytest.raises(api.WorkbookError, match="no cached results"):
        api.import_workbook(filled_workbook(tmp_path, cached=False), client_, tender, today=TODAY)
    (tmp_path / "other.xlsx").write_bytes(b"not a workbook")
    with pytest.raises(api.WorkbookError, match="not a readable"):
        api.import_workbook(tmp_path / "other.xlsx", client_, tender)


def test_missing_resource_is_reported(tender, tmp_path):
    bare = Client.objects.create(name="Πελάτης Δύο", tax_id="111111114")
    with pytest.raises(api.Differences) as refused:
        api.import_workbook(filled_workbook(tmp_path), bare, tender, today=TODAY)
    assert "no resource labelled 'όχημα 1'" in refused.value.report.text()


def test_command(client_, tmp_path, capsys):
    path = filled_workbook(tmp_path)
    call_command("import_v0", str(path), "--client", str(client_.pk), "--tender", "pkm-meth-student-transport-dsa-2026")
    assert "differences: 0" in capsys.readouterr().out
    with pytest.raises(CommandError, match="no client"):
        call_command("import_v0", str(path), "--client", "999999", "--tender", "pkm-meth-student-transport-dsa-2026")


def test_a_guarantee_below_the_minimum_fails_r8(client_, tender, tmp_path):
    path = filled_workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Προσφορά"]["E26"] = wb["Προσφορά"]["E26"].value - 0.01  # one cent short
    wb["Προσφορά"]["E29"] = "ΛΑΘΟΣ: ποσό κάτω από το ελάχιστο (§4.3.1.2)"
    wb["Έγγραφα"]["J7"] = "NOT_SATISFIED"
    wb.save(path)
    report = api.import_workbook(path, client_, tender, today=TODAY)
    assert report.differences == 0, report.text()
    r8 = next(i for i in Bid.objects.get().checklist_snapshot if i["requirement_id"] == "R8")
    assert (r8["status"], r8["reason"]) == ("NOT_SATISFIED", "guarantee_below_minimum")
